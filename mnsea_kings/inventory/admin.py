from io import BytesIO
from datetime import timedelta
from copy import copy
from pathlib import Path

from django.contrib import admin, messages
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import path
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Font, PatternFill

from catalog.models import Product

from .models import StockDay, StockLine, StockMonth, StockTransaction


class StockLineInline(admin.TabularInline):
    model = StockLine
    extra = 1
    fields = (
        "product", "opening_qty", "received_qty", "sold_qty",
        "returns_qty", "wastage_qty", "automatic_sold_display", "closing_display", "remarks",
    )
    readonly_fields = ("opening_qty", "automatic_sold_display", "closing_display")

    @admin.display(description="Auto sold")
    def automatic_sold_display(self, obj):
        return obj.automatic_sold_qty if obj.pk else "-"

    @admin.display(description="Closing")
    def closing_display(self, obj):
        return obj.closing_qty if obj.pk else "-"


@admin.register(StockDay)
class StockDayAdmin(admin.ModelAdmin):
    list_display = ("date", "stock_month", "line_count", "day_closing_total")
    list_filter = ("stock_month",)
    date_hierarchy = "date"
    search_fields = ("notes",)
    inlines = [StockLineInline]

    @admin.display(description="Products")
    def line_count(self, obj):
        return obj.lines.count()

    @admin.display(description="Closing quantity")
    def day_closing_total(self, obj):
        return sum((line.closing_qty for line in obj.lines.all()), 0)


@admin.action(description="Export selected month as Excel workbook")
def export_stock_book(modeladmin, request, queryset):
    book = queryset.order_by("year", "month").first()
    if not book:
        modeladmin.message_user(request, "Select one monthly stock book.", messages.WARNING)
        return
    template_path = Path(__file__).resolve().parent.parent / "static" / "images" / "KNL_SIH_Stock_Sheet.xlsx"
    workbook = load_workbook(template_path)
    template = workbook.active
    template.title = "01-Template"
    summary = workbook.create_sheet("Monthly Summary", 0)
    _format_sample_sheet(summary, f"{book} | Monthly Summary")
    summary_headers = ["Date", "Sku", "QTY", "Returns Qty", "Total Qty", "Remarks"]
    for column, value in enumerate(summary_headers, 1):
        summary.cell(2, column, value)
    days = book.days.prefetch_related("lines__product__category").order_by("date")
    for index, day in enumerate(days):
        sheet = template if index == 0 else workbook.copy_worksheet(template)
        sheet.title = day.date.strftime("%d-%b")[:31]
        _format_sample_sheet(sheet, f"{book} | {day.date:%A, %d %B %Y}")
        lines = list(day.lines.all())
        for row, line in enumerate(lines, 3):
            available_qty = line.opening_qty + line.received_qty - line.sold_qty - line.wastage_qty
            sheet.cell(row, 1, line.product.name)
            sheet.cell(row, 2, available_qty)
            sheet.cell(row, 3, line.returns_qty)
            sheet.cell(row, 4, f"=SUM(B{row}:C{row})")
            sheet.cell(row, 5, "")
            sheet.cell(row, 6, line.remarks)
            _style_sample_row(sheet, row, alternate=row % 2 == 0)
            summary.append([day.date, line.product.name, available_qty, line.returns_qty, f"=SUM(C{summary.max_row + 1}:D{summary.max_row + 1})", line.remarks])
        total_row = len(lines) + 3
        _style_total_row(sheet, total_row)
        sheet.cell(total_row, 1, "TOTAL")
        for column in (2, 3, 4):
            letter = chr(64 + column)
            sheet.cell(total_row, column, f"=SUM({letter}3:{letter}{total_row - 1})")
        note_row = total_row + 2
        sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
        sheet.cell(note_row, 1, "Note: QTY excludes returns; Total Qty includes returns. Returns are carried into the next day's opening stock.")
        sheet.cell(note_row, 1)._style = copy(sheet[37][0]._style)
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:F{total_row}"
        if index == 0:
            template = sheet
    _format_summary_sheet(summary, summary_headers)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="stock-book-{book.year}-{book.month:02d}.xlsx"'
    return response


def _format_sample_sheet(sheet, title):
    sheet.delete_rows(1, sheet.max_row)
    sheet.merge_cells("A1:F1")
    sheet["A1"] = title
    sheet["A1"]._style = copy(sheet["A1"]._style)
    sheet["A1"].font = Font(name="Arial", size=13, bold=True, color="FF000000")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    sheet["A1"].alignment = copy(sheet["A1"].alignment)
    headers = ["Sku", "QTY", "Returns Qty", "Total Qty", "Aging", "Crate Numbers"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(2, column, value)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF" if column == 1 else "FF1F1F1F")
        cell.fill = PatternFill("solid", fgColor="FF2E75B6" if column == 1 else "FFBDD7EE")
        cell.alignment = copy(sheet["A2"].alignment)
    for column, width in {"A": 40, "B": 11, "C": 13, "D": 12, "E": 10, "F": 15}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"


def _style_sample_row(sheet, row, alternate=False):
    fill = PatternFill("solid", fgColor="FFF2F2F2" if alternate else "FFFFFFFF")
    for column in range(1, 7):
        cell = sheet.cell(row, column)
        cell.fill = fill
        cell.font = Font(name="Arial", size=10.5, color="FF1F4E78" if column in (2, 3, 5, 6) else "FF000000", bold=column == 4)
        cell.alignment = copy(sheet["A3"].alignment)


def _style_total_row(sheet, row):
    for column in range(1, 7):
        cell = sheet.cell(row, column)
        cell.fill = PatternFill("solid", fgColor="FF92D050")
        cell.font = Font(name="Arial", size=11, bold=True, color="FF000000")
        cell.alignment = copy(sheet["A2"].alignment)


def _format_summary_sheet(sheet, headers):
    for column, value in enumerate(headers, 1):
        sheet.cell(2, column, value)
    for row in range(3, sheet.max_row + 1):
        _style_sample_row(sheet, row, alternate=row % 2 == 0)
    total_row = sheet.max_row + 1
    _style_total_row(sheet, total_row)
    sheet.cell(total_row, 1, "TOTAL")
    for column in (3, 4, 5):
        letter = chr(64 + column)
        sheet.cell(total_row, column, f"=SUM({letter}3:{letter}{total_row - 1})")
    sheet.auto_filter.ref = f"A2:F{total_row}"


@admin.action(description="Close selected month (prevents edits)")
def close_stock_book(modeladmin, request, queryset):
    approved = queryset.filter(is_approved=True, is_closed=False)
    updated = approved.update(is_closed=True, closed_at=timezone.now())
    modeladmin.message_user(request, f"Closed {updated} stock book(s).", messages.SUCCESS)
    if queryset.filter(is_approved=False, is_closed=False).exists():
        modeladmin.message_user(request, "Some books were not closed because they need approval first.", messages.WARNING)


@admin.action(description="Approve selected month")
def approve_stock_book(modeladmin, request, queryset):
    approved = 0
    for book in queryset.filter(is_closed=False, is_approved=False):
        book.approve(request.user)
        approved += 1
    modeladmin.message_user(request, f"Approved {approved} stock book(s).", messages.SUCCESS)


@admin.action(description="Generate all daily tabs and active products")
def generate_stock_days(modeladmin, request, queryset):
    created_days = 0
    created_lines = 0
    products = Product.objects.filter(is_available=True).select_related("category")
    for book in queryset.filter(is_closed=False):
        current = book.first_day
        while current < book.last_day:
            day, day_created = StockDay.objects.get_or_create(stock_month=book, date=current)
            created_days += int(day_created)
            for product in products:
                _, line_created = StockLine.objects.get_or_create(day=day, product=product)
                created_lines += int(line_created)
            current += timedelta(days=1)
    modeladmin.message_user(
        request, f"Generated {created_days} day(s) and {created_lines} product line(s).", messages.SUCCESS
    )


@admin.action(description="Recalculate carried-forward opening quantities")
def recalculate_openings(modeladmin, request, queryset):
    changed = 0
    for book in queryset.filter(is_closed=False):
        for line in StockLine.objects.filter(day__stock_month=book).select_related("day", "day__stock_month"):
            opening = line.carried_forward_opening()
            if line.opening_qty != opening:
                StockLine.objects.filter(pk=line.pk).update(opening_qty=opening)
                changed += 1
    modeladmin.message_user(request, f"Recalculated {changed} opening balance(s).", messages.SUCCESS)


@admin.register(StockMonth)
class StockMonthAdmin(admin.ModelAdmin):
    list_display = ("__str__", "year", "month", "day_count", "is_approved", "is_closed", "closed_at")
    list_filter = ("year", "month", "is_closed")
    actions = (generate_stock_days, recalculate_openings, approve_stock_book, export_stock_book, close_stock_book)
    ordering = ("-year", "-month")

    def get_urls(self):
        urls = super().get_urls()
        custom = [
            path("dashboard/", self.admin_site.admin_view(stock_dashboard), name="inventory_stock_dashboard"),
            path("import/", self.admin_site.admin_view(import_stock_book), name="inventory_stock_import"),
        ]
        return custom + urls

    @admin.display(description="Days entered")
    def day_count(self, obj):
        return obj.days.count()


@admin.register(StockLine)
class StockLineAdmin(admin.ModelAdmin):
    list_display = (
        "day", "category_name", "product", "opening_qty", "received_qty",
        "sold_qty", "returns_qty", "wastage_qty", "closing_display",
    )
    list_filter = ("day__stock_month", "product__category")
    search_fields = ("product__name", "product__category__name")
    readonly_fields = ("opening_qty", "closing_display")

    @admin.display(description="Closing")
    def closing_display(self, obj):
        return obj.closing_qty


@admin.register(StockTransaction)
class StockTransactionAdmin(admin.ModelAdmin):
    list_display = (
        "transaction_date", "product", "transaction_type", "quantity",
        "signed_quantity", "order", "created_by", "created_at",
    )
    list_filter = ("transaction_type", "transaction_date", "product__category")
    search_fields = ("product__name", "reason", "order__order_number")
    readonly_fields = ("created_at", "source_key", "order")
    date_hierarchy = "transaction_date"

    def save_model(self, request, obj, form, change):
        if not change:
            obj.created_by = request.user
        super().save_model(request, obj, form, change)


def _dashboard_data():
    today = timezone.localdate()
    transactions = StockTransaction.objects.filter(transaction_date=today)
    stock_books = StockMonth.objects.select_related("approved_by").order_by("-year", "-month")
    latest_book = stock_books.first()
    approved_book = stock_books.filter(is_approved=True).first()
    approved_lines = list(
        StockLine.objects.filter(day__stock_month=approved_book).select_related("product__category")
    ) if approved_book else []
    approved_category_totals = {}
    for line in approved_lines:
        category = line.product.category.name
        approved_category_totals[category] = approved_category_totals.get(category, 0) + line.closing_qty
    sales = transactions.filter(transaction_type="sale").aggregate(total=Sum("quantity"))["total"] or 0
    returns = transactions.filter(transaction_type="customer_return").aggregate(total=Sum("quantity"))["total"] or 0
    wastage = transactions.filter(transaction_type="wastage").aggregate(total=Sum("quantity"))["total"] or 0
    received = transactions.filter(transaction_type="purchase").aggregate(total=Sum("quantity"))["total"] or 0
    return {
        "today": today,
        "sales": sales,
        "returns": returns,
        "wastage": wastage,
        "received": received,
        "movement_total": sales + returns + wastage + received,
        "latest_book": latest_book,
        "open_books": stock_books.filter(is_closed=False).count(),
        "approved_books": stock_books.filter(is_approved=True).count(),
        "closed_books": stock_books.filter(is_closed=True).count(),
        "approved_book": approved_book,
        "approved_received": sum((line.received_qty for line in approved_lines), 0),
        "approved_returns": sum((line.effective_returns_qty for line in approved_lines), 0),
        "approved_closing": sum((line.closing_qty for line in approved_lines), 0),
        "approved_category_totals": sorted(approved_category_totals.items()),
        "category_totals": transactions.values("product__category__name").annotate(total=Sum("quantity")).order_by("product__category__name"),
        "low_stock": [line for line in StockLine.objects.filter(day__date=today).select_related("product") if line.closing_qty <= 0],
    }


def stock_dashboard(request):
    return render(request, "admin/inventory/dashboard.html", _dashboard_data())


@transaction.atomic
def import_stock_book(request):
    if request.method == "POST":
        upload = request.FILES.get("workbook")
        year = request.POST.get("year")
        month = request.POST.get("month")
        if not upload or not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "Upload an .xlsx workbook.")
        else:
            try:
                workbook = load_workbook(upload, data_only=False)
                if not year or not month or not (str(year).isdigit() and str(month).isdigit() and 1 <= int(month) <= 12):
                    raise ValueError("Enter a valid year and month.")
                from catalog.models import Product
                from catalog.models import Category
                from datetime import datetime
                book, _ = StockMonth.objects.get_or_create(year=int(year), month=int(month))
                if book.is_closed:
                    raise ValueError("The selected stock month is closed.")
                product_map = {product.name.casefold(): product for product in Product.objects.all()}
                imported = 0
                duplicates = 0
                created_products = 0
                recognized_sheets = 0
                for sheet in workbook.worksheets:
                    if sheet.title in ("Monthly Summary", "01-Template"):
                        continue
                    try:
                        day_date = datetime.strptime(f"{sheet.title}-{year}", "%d-%b-%Y").date()
                    except ValueError:
                        title_value = str(sheet.cell(1, 1).value or "")
                        try:
                            day_date = datetime.strptime(title_value.rsplit("-", 1)[-1].strip(), "%d/%m/%Y").date()
                        except ValueError:
                            continue
                    if not (book.first_day <= day_date < book.last_day):
                        continue
                    recognized_sheets += 1
                    day, _ = StockDay.objects.get_or_create(stock_month=book, date=day_date)
                    for row in sheet.iter_rows(min_row=3, values_only=True):
                        if not row[0] or str(row[0]).upper() == "TOTAL":
                            continue
                        product = product_map.get(str(row[0]).strip().casefold())
                        if not product:
                            imported_category, _ = Category.objects.get_or_create(
                                slug="imported-stock",
                                defaults={"name": "Imported Stock", "is_active": False},
                            )
                            product, product_created = Product.objects.get_or_create(
                                name=str(row[0]).strip(),
                                defaults={
                                    "category": imported_category,
                                    "price": 0,
                                    "unit": "kg",
                                    "is_available": False,
                                },
                            )
                            product_map[product.name.casefold()] = product
                            created_products += int(product_created)
                        qty = row[1] or 0
                        returns = row[2] or 0
                        if not isinstance(qty, (int, float)) or not isinstance(returns, (int, float)):
                            raise ValueError(f"Invalid quantity for {row[0]} on {sheet.title}.")
                        line, created = StockLine.objects.get_or_create(day=day, product=product)
                        if not created and (line.received_qty or line.sold_qty or line.returns_qty):
                            duplicates += 1
                            continue
                        line.received_qty = qty
                        line.returns_qty = returns
                        line.remarks = "Imported from workbook"
                        line.save()
                        imported += 1
                if not recognized_sheets:
                    raise ValueError("No dated sheets found. Use daily tabs such as 01-Aug or put the date in A1 as DD/MM/YYYY.")
                if not imported and duplicates:
                    messages.warning(request, f"No changes made: all {duplicates} matching rows already contain data.")
                elif not imported:
                    raise ValueError("No product rows were imported. Check that Sku names exactly match Catalog products.")
                else:
                    messages.success(request, f"Imported {imported} row(s); created {created_products} inactive product(s); skipped {duplicates} duplicate row(s).")
            except (ValueError, TypeError, KeyError, OverflowError, InvalidFileException) as error:
                transaction.set_rollback(True)
                messages.error(request, f"Workbook could not be read: {error}")
        return redirect(request.path)
    return render(request, "admin/inventory/import.html")
